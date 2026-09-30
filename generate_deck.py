"""
generate_deck.py - Generates PPTX and interactive HTML presentations for LearnMate AI.
Team Nova Stack | Prompt to Production Challenge
"""
import os
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_DIR = Path(__file__).parent
PPTX_FILE = OUTPUT_DIR / "learnmate_ai_presentation.pptx"
HTML_FILE = OUTPUT_DIR / "presentation.html"

# Slide Data Definition
SLIDES_DATA = [
    {
        "slide_num": 1,
        "category": "TITLE & VISION",
        "title": "LearnMate AI",
        "subtitle": "The Adaptive Micro-Sprint Mentor",
        "team": "Team Nova Stack | Prompt to Production Challenge",
        "bullets": [
            "Time-boxed adaptive mentor tailored to your schedule.",
            "Measures blind spots and diagnoses cognitive misconceptions.",
            "Rebuilds future lessons to cure learning gaps.",
        ],
        "visual_label": "✦ System Core Concept",
        "visual_items": [
            ("⏱️ Time-Boxed", "Guaranteed completion within 6 to 60 minutes"),
            ("🧠 Adaptive Engine", "Detects weak concepts (< 0.6) and adjusts next lesson"),
            ("🗣️ Multilingual", "English, Telugu, Hindi vernacular concept explanations"),
        ],
        "speaker_note": "Hello judges! We are team Nova Stack, presenting LearnMate AI. Most online learning treats education like a firehose. LearnMate AI flips the script—it teaches strictly within your available time budget, pinpoints exactly what you didn't understand, and reconstructs the next lesson to fix your cognitive blind spots.",
    },
    {
        "slide_num": 2,
        "category": "THE PROBLEM",
        "title": "Why Self-Learning Breaks",
        "subtitle": "The Chasm Between Content and Retention",
        "team": "Team Nova Stack | Prompt to Production Challenge",
        "bullets": [
            "Static courses ignore learners' actual available minutes.",
            "Generic LLM chatbots forget previous student mistakes.",
            "Passive reading fails to detect cognitive blindspots.",
        ],
        "visual_label": "⚠️ The Status Quo Breakdown",
        "visual_items": [
            ("❌ 4-Hour Videos", "Learners run out of time and abandon courses mid-way"),
            ("❌ Amnesic Chatbots", "No memory of previous mistakes or structured progress"),
            ("❌ Hidden Misconceptions", "Silent failure on foundational syntax & boundary rules"),
        ],
        "speaker_note": "Modern learners have limited windows to study, yet existing platforms offer rigid hour-long tutorials that ignore your schedule. On the other hand, generic AI chatbots are completely amnesic—they don't measure mastery or remember past errors. Students fall through the cracks without targeted feedback.",
    },
    {
        "slide_num": 3,
        "category": "THE SOLUTION",
        "title": "Time-Boxed, Personalized Mentorship",
        "subtitle": "Learner Agency with Algorithmic Precision",
        "team": "Team Nova Stack | Prompt to Production Challenge",
        "bullets": [
            "Learner inputs topic, minutes, level, and language.",
            "Multilingual mentoring across English, Telugu, and Hindi.",
            "Strictly time-boxed 6-section personalized curriculum.",
        ],
        "visual_label": "🎯 Learner Configuration Interface",
        "visual_items": [
            ("✨ Any Topic Input", "Free text + Quick chips (Python loops, SQL joins, Photosynthesis)"),
            ("⏱️ Precision Budget", "Portions 6 sections (Intro, Explain, Example, Practice, Quiz, Recap)"),
            ("🌐 Regional Support", "Authentic vernacular cues while code syntax stays in English"),
        ],
        "speaker_note": "LearnMate AI solves this by giving complete agency back to the learner. You type any topic, choose your time budget from 6 to 60 minutes, and select English, Telugu, or Hindi. LearnMate immediately portions out and synthesizes a high-yield, 6-section lesson guaranteed to respect your schedule.",
    },
    {
        "slide_num": 4,
        "category": "THE ADAPTIVE LOOP",
        "title": "Code Decides. AI Explains.",
        "subtitle": "The Deterministic Pedagogical State Machine",
        "team": "Team Nova Stack | Prompt to Production Challenge",
        "bullets": [
            "Deterministic code drives decisions; AI writes explanations.",
            "Closed loop: Lesson, Quiz, Mastery, then Remediation.",
            "Halts progression until misconceptions are actively cured.",
        ],
        "visual_label": "🔄 The Continuous Quality Flywheel",
        "visual_items": [
            ("1. Time-Boxed Lesson", "Portioned sections with embedded interactive micro-simulators"),
            ("2. 5-Question Checkpoint", "Targeted quiz items tied strictly to evaluated concept keys"),
            ("3. SQLite Mastery EMA", "Calculates mastery: mastery += 0.3 * (outcome - mastery)"),
            ("4. Adaptive Sprint", "If mastery < 0.6, next lesson restructures around blindspot"),
        ],
        "speaker_note": "Here is our core architectural philosophy: deterministic code makes pedagogical decisions, while Gemini provides rich explanations. After every sprint, a diagnostic quiz evaluates understanding. If concept mastery dips below 0.6, our engine catches it and restructures the very next lesson around that misconception.",
    },
    {
        "slide_num": 5,
        "category": "SYSTEM ARCHITECTURE",
        "title": "Production-Grade Reliability",
        "subtitle": "Engineered for Resilience, Speed, and Zero Crashes",
        "team": "Team Nova Stack | Prompt to Production Challenge",
        "bullets": [
            "Streamlit UI featuring interactive execution visualizers.",
            "Gemini JSON mode validated with strict Pydantic schemas.",
            "SQLite EMA tracker calculating concept mastery updates.",
        ],
        "visual_label": "🏗️ Full Stack Architecture",
        "visual_items": [
            ("🖥️ Streamlit Frontend", "SaaS design, dark theme, Range Visualizer & Step Stepper"),
            ("🤖 Gemini 1.5 Flash", "JSON mode, Pydantic validation, single retry & graceful fallback"),
            ("🗄️ SQLite Engine", "Compound primary key (topic, concept_id) preventing cross-topic leaks"),
            ("⚡ Local Disk Cache", "Cache by topic+level+language+minutes+weak_concept in cache/"),
        ],
        "speaker_note": "Under the hood, LearnMate is engineered for resilience. We combine a responsive Streamlit UI with interactive canvas micro-simulators. Gemini runs in JSON mode validated strictly through Pydantic with automatic retry guards, while SQLite maintains concept mastery using an exponential moving average formula.",
    },
    {
        "slide_num": 6,
        "category": "DEMO IN ACTION",
        "title": "Diagnosing & Curing Mistakes",
        "subtitle": "From Cognitive Slip to Permanent Mastery",
        "team": "Team Nova Stack | Prompt to Production Challenge",
        "bullets": [
            "Student selects Python loops; misses exclusive range bound.",
            "Engine isolates range_bounds weakness below threshold.",
            "Remedial lesson injects 'Your Mistake' box and traces.",
        ],
        "visual_label": "🔬 Live Diagnostic Flow",
        "visual_items": [
            ("1. Checkpoint Miss", "User assumes range(0, 5) yields 5. Outcome = 0.0"),
            ("2. Mastery Plunge", "Mastery drops: 0.5 -> 0.35. Weak concept diagnosed: range_bounds"),
            ("3. 'Your Mistake' Card", "Explains exact error, student choice vs correct answer"),
            ("4. Execution Trace", "Step-by-step table + Stop boundary visualizer halt simulation"),
        ],
        "speaker_note": "In our live demo, a student takes a 15-minute Python sprint and assumes range(0, 5) includes 5. The engine instantly isolates range_bounds. The adaptive lesson then deploys a dedicated 'Your Mistake' breakdown, a step-by-step trace table, and an interactive stepper. And this works dynamically for any topic typed.",
    },
    {
        "slide_num": 7,
        "category": "IMPACT & ROADMAP",
        "title": "What's Next for LearnMate",
        "subtitle": "Democratizing 1-on-1 Adaptive Tutoring",
        "team": "Team Nova Stack | Prompt to Production Challenge",
        "bullets": [
            "Expanded interactive visualizers and additional vernacular languages.",
            "Teacher analytics dashboard for classroom-wide blindspot tracking.",
            "Hybrid offline edge-model inference with cached fallbacks.",
        ],
        "visual_label": "🚀 Roadmap & Live Demo",
        "visual_items": [
            ("🎨 Rich Visual Steppers", "Expanding generic step_viz to dynamic chemistry & history diagrams"),
            ("📊 Classroom Dashboard", "Aggregated SQLite mastery metrics for school instructors"),
            ("📱 Offline On-Device", "Ultra-fast local model execution with zero cloud dependency"),
            ("✦ Live Demo Ready", "Thank You! Experience LearnMate AI live now"),
        ],
        "speaker_note": "LearnMate bridges the gap between passive video streaming and expensive private tutoring. Looking ahead, we're expanding our generic step-visualizer components, onboarding more Indian regional languages, and building a classroom dashboard for teachers. Thank you, and try our live demo today!",
    },
]


def create_pptx_deck():
    """Builds a 16:9 widescreen PowerPoint deck with dark modern aesthetics."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    c_bg = RGBColor(11, 15, 25)        # #0b0f19
    c_card = RGBColor(17, 24, 39)      # #111827
    c_card_border = RGBColor(31, 41, 55) # #1f2937
    c_cyan = RGBColor(56, 189, 248)    # #38bdf8
    c_indigo = RGBColor(129, 140, 248) # #818cf8
    c_white = RGBColor(248, 250, 252)  # #f8fafc
    c_gray = RGBColor(148, 163, 184)   # #94a3b8
    c_emerald = RGBColor(52, 211, 153) # #34d399

    for s_idx, data in enumerate(SLIDES_DATA):
        slide = prs.slides.add_slide(blank_layout)

        # 1. Background Fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = c_bg
        bg_shape.line.fill.background()

        # 2. Top Header Bar: Category pill + Team watermark
        # Category Badge
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(5.0), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = f"✦ {data['category']}  |  SLIDE {data['slide_num']} OF 7"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_cyan

        # Team watermark (right)
        team_box = slide.shapes.add_textbox(Inches(7.5), Inches(0.45), Inches(5.0), Inches(0.4))
        tf_team = team_box.text_frame
        tf_team.word_wrap = True
        p_team = tf_team.paragraphs[0]
        p_team.alignment = PP_ALIGN.RIGHT
        p_team.text = "TEAM NOVA STACK  •  PROMPT TO PRODUCTION"
        p_team.font.size = Pt(10)
        p_team.font.bold = True
        p_team.font.color.rgb = c_gray

        # 3. Main Title & Subtitle
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(1.3))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(36)
        p_title.font.bold = True
        p_title.font.color.rgb = c_white

        p_sub = tf_title.add_paragraph()
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = c_cyan
        p_sub.space_before = Pt(4)

        # 4. Left Card: The 3 Core Bullets (max 10 words each)
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(5.7), Inches(4.5))
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = c_card
        left_card.line.color.rgb = c_card_border
        left_card.line.width = Pt(1.5)

        # Content in Left Card
        left_tb = slide.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(5.3), Inches(4.1))
        tf_left = left_tb.text_frame
        tf_left.word_wrap = True
        
        p_lh = tf_left.paragraphs[0]
        p_lh.text = "KEY TAKEAWAYS"
        p_lh.font.size = Pt(11)
        p_lh.font.bold = True
        p_lh.font.color.rgb = c_indigo

        for bullet in data["bullets"]:
            p_b = tf_left.add_paragraph()
            p_b.text = f"•  {bullet}"
            p_b.font.size = Pt(15.5)
            p_b.font.color.rgb = c_white
            p_b.space_before = Pt(20)
            p_b.line_spacing = 1.25

        # 5. Right Card: Visual Component / Architecture / Demo Highlight
        right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.4), Inches(5.7), Inches(4.5))
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = c_card
        right_card.line.color.rgb = c_cyan
        right_card.line.width = Pt(1.5)

        right_tb = slide.shapes.add_textbox(Inches(7.0), Inches(2.6), Inches(5.3), Inches(4.1))
        tf_right = right_tb.text_frame
        tf_right.word_wrap = True

        p_rh = tf_right.paragraphs[0]
        p_rh.text = data["visual_label"]
        p_rh.font.size = Pt(11)
        p_rh.font.bold = True
        p_rh.font.color.rgb = c_cyan

        for v_head, v_desc in data["visual_items"]:
            p_vh = tf_right.add_paragraph()
            p_vh.text = v_head
            p_vh.font.size = Pt(14)
            p_vh.font.bold = True
            p_vh.font.color.rgb = c_emerald
            p_vh.space_before = Pt(12)

            p_vd = tf_right.add_paragraph()
            p_vd.text = v_desc
            p_vd.font.size = Pt(12)
            p_vd.font.color.rgb = c_gray
            p_vd.space_before = Pt(2)

        # 6. Speaker Notes
        notes_slide = slide.notes_slide
        tf_notes = notes_slide.notes_text_frame
        tf_notes.text = f"SPEAKER NOTE (20 SECONDS):\n\n{data['speaker_note']}"

    prs.save(str(PPTX_FILE))
    print(f"✅ Generated PowerPoint Presentation: {PPTX_FILE}")


def create_html_presentation():
    """Generates an interactive, dark, modern HTML slide deck with presenter controls."""
    slides_json = json.dumps(SLIDES_DATA)
    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>LearnMate AI — Hackathon Pitch Deck | Team Nova Stack</title>
  <style>
    * {{
      box-sizing: border-box;
      margin: 0;
      padding: 0;
      font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
    }}
    body {{
      background: #070a12;
      color: #f8fafc;
      overflow: hidden;
      height: 100vh;
      display: flex;
      flex-direction: column;
    }}
    /* Top Bar */
    header {{
      background: rgba(11, 15, 25, 0.95);
      border-bottom: 1px solid rgba(255, 255, 255, 0.08);
      padding: 12px 28px;
      display: flex;
      justify-content: space-between;
      align-items: center;
      z-index: 10;
    }}
    .brand {{
      display: flex;
      align-items: center;
      gap: 10px;
      font-weight: 800;
      font-size: 17px;
      color: #38bdf8;
      letter-spacing: -0.5px;
    }}
    .brand-tag {{
      background: rgba(56, 189, 248, 0.15);
      border: 1px solid rgba(56, 189, 248, 0.3);
      color: #38bdf8;
      font-size: 11px;
      font-weight: 700;
      padding: 3px 8px;
      border-radius: 999px;
    }}
    .meta-bar {{
      display: flex;
      gap: 16px;
      align-items: center;
      font-size: 13px;
      color: #94a3b8;
    }}
    .team-badge {{
      color: #a5b4fc;
      font-weight: 700;
    }}

    /* Slide Viewport */
    main {{
      flex: 1;
      display: flex;
      justify-content: center;
      align-items: center;
      padding: 24px 36px;
      position: relative;
    }}
    .slide-stage {{
      width: 100%;
      max-width: 1240px;
      height: 100%;
      max-height: 700px;
      background: #0b0f19;
      border: 1px solid rgba(255, 255, 255, 0.1);
      border-radius: 20px;
      padding: 40px 50px;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      box-shadow: 0 25px 60px -15px rgba(0, 0, 0, 0.8), 0 0 40px -10px rgba(56, 189, 248, 0.15);
      position: relative;
      transition: all 0.3s ease;
    }}

    .slide-category {{
      display: inline-flex;
      align-items: center;
      gap: 6px;
      font-size: 12px;
      font-weight: 800;
      color: #38bdf8;
      letter-spacing: 1.5px;
      text-transform: uppercase;
      margin-bottom: 6px;
    }}
    .slide-title {{
      font-size: 42px;
      font-weight: 800;
      letter-spacing: -1px;
      background: linear-gradient(135deg, #ffffff 0%, #cbd5e1 100%);
      -webkit-background-clip: text;
      -webkit-text-fill-color: transparent;
      margin-bottom: 4px;
    }}
    .slide-subtitle {{
      font-size: 18px;
      color: #38bdf8;
      font-weight: 500;
      margin-bottom: 28px;
    }}

    /* 2-Column Content Layout */
    .slide-body {{
      display: grid;
      grid-template-columns: 1.15fr 1fr;
      gap: 28px;
      flex: 1;
      min-height: 0;
    }}
    .card {{
      background: #111827;
      border: 1px solid #1f2937;
      border-radius: 14px;
      padding: 24px;
      display: flex;
      flex-direction: column;
      justify-content: flex-start;
    }}
    .card-highlight {{
      border-color: rgba(56, 189, 248, 0.4);
      background: linear-gradient(145deg, #111827 0%, rgba(15, 23, 42, 0.8) 100%);
    }}
    .card-title {{
      font-size: 12px;
      font-weight: 800;
      letter-spacing: 1px;
      color: #818cf8;
      text-transform: uppercase;
      margin-bottom: 18px;
    }}
    .card-highlight .card-title {{
      color: #38bdf8;
    }}

    .bullet-list {{
      list-style: none;
      display: flex;
      flex-direction: column;
      gap: 18px;
    }}
    .bullet-item {{
      display: flex;
      align-items: flex-start;
      gap: 12px;
      font-size: 19px;
      font-weight: 600;
      color: #f1f5f9;
      line-height: 1.4;
    }}
    .bullet-dot {{
      width: 8px;
      height: 8px;
      border-radius: 50%;
      background: #38bdf8;
      margin-top: 10px;
      flex-shrink: 0;
      box-shadow: 0 0 10px #38bdf8;
    }}

    .visual-list {{
      display: flex;
      flex-direction: column;
      gap: 12px;
    }}
    .visual-node {{
      background: #090d16;
      border: 1px solid rgba(255, 255, 255, 0.06);
      border-radius: 10px;
      padding: 12px 14px;
    }}
    .visual-head {{
      font-size: 14.5px;
      font-weight: 700;
      color: #34d399;
      margin-bottom: 3px;
    }}
    .visual-desc {{
      font-size: 13px;
      color: #94a3b8;
      line-height: 1.4;
    }}

    /* Notes Drawer */
    .notes-drawer {{
      background: #0f172a;
      border: 1px solid #334155;
      border-radius: 12px;
      padding: 14px 18px;
      margin-top: 20px;
      display: none;
    }}
    .notes-header {{
      display: flex;
      align-items: center;
      gap: 8px;
      font-size: 11px;
      font-weight: 800;
      color: #fbbf24;
      text-transform: uppercase;
      margin-bottom: 6px;
    }}
    .notes-text {{
      font-size: 14px;
      color: #e2e8f0;
      line-height: 1.5;
      font-style: italic;
    }}

    /* Footer Controls */
    footer {{
      background: rgba(11, 15, 25, 0.95);
      border-top: 1px solid rgba(255, 255, 255, 0.08);
      padding: 14px 36px;
      display: flex;
      justify-content: space-between;
      align-items: center;
    }}
    .progress-wrap {{
      display: flex;
      align-items: center;
      gap: 14px;
    }}
    .progress-bar {{
      width: 240px;
      height: 6px;
      background: #1e293b;
      border-radius: 999px;
      overflow: hidden;
    }}
    .progress-fill {{
      height: 100%;
      background: linear-gradient(90deg, #38bdf8 0%, #818cf8 100%);
      transition: width 0.3s ease;
    }}
    .step-counter {{
      font-size: 13px;
      font-weight: 700;
      color: #cbd5e1;
    }}
    .controls {{
      display: flex;
      gap: 12px;
      align-items: center;
    }}
    button {{
      background: #1e293b;
      border: 1px solid #334155;
      color: #f8fafc;
      padding: 8px 18px;
      border-radius: 8px;
      font-weight: 600;
      font-size: 13px;
      cursor: pointer;
      display: flex;
      align-items: center;
      gap: 6px;
      transition: all 0.2s;
    }}
    button:hover:not(:disabled) {{
      background: #0284c7;
      border-color: #38bdf8;
      color: white;
      transform: translateY(-1px);
    }}
    button:disabled {{
      opacity: 0.4;
      cursor: not-allowed;
    }}
    .btn-notes {{
      background: rgba(251, 191, 36, 0.15);
      border-color: rgba(251, 191, 36, 0.3);
      color: #fbbf24;
    }}
    .btn-notes:hover {{
      background: #fbbf24;
      color: #0b0f19;
    }}
  </style>
</head>
<body>

  <header>
    <div class="brand">
      <span>✦ LearnMate AI</span>
      <span class="brand-tag">Prompt to Production</span>
    </div>
    <div class="meta-bar">
      <span class="team-badge">Team Nova Stack</span>
      <span>•</span>
      <span>7-Slide Hackathon Pitch</span>
    </div>
  </header>

  <main>
    <div class="slide-stage" id="stage">
      <div>
        <div class="slide-category" id="cat-label">✦ TITLE &amp; VISION</div>
        <h1 class="slide-title" id="slide-title">LearnMate AI</h1>
        <div class="slide-subtitle" id="slide-sub">The Adaptive Micro-Sprint Mentor</div>
      </div>

      <div class="slide-body">
        <div class="card">
          <div class="card-title">Key Value Proposition</div>
          <ul class="bullet-list" id="bullet-container"></ul>
        </div>
        <div class="card card-highlight">
          <div class="card-title" id="visual-title">System Highlight</div>
          <div class="visual-list" id="visual-container"></div>
        </div>
      </div>

      <div class="notes-drawer" id="notes-drawer">
        <div class="notes-header">🎙️ 20-Second Speaker Pitch Note</div>
        <p class="notes-text" id="notes-content"></p>
      </div>
    </div>
  </main>

  <footer>
    <div class="progress-wrap">
      <div class="progress-bar">
        <div class="progress-fill" id="prog-fill" style="width: 14%;"></div>
      </div>
      <span class="step-counter" id="counter-label">Slide 1 of 7</span>
    </div>

    <div class="controls">
      <button class="btn-notes" onclick="toggleNotes()">🎙️ Speaker Notes (20s)</button>
      <button id="btn-prev" onclick="changeSlide(-1)" disabled>⬅ Prev</button>
      <button id="btn-next" onclick="changeSlide(1)">Next ➡</button>
    </div>
  </footer>

  <script>
    const deck = {slides_json};
    let currentIdx = 0;
    let showNotes = true;

    function renderSlide() {{
      const data = deck[currentIdx];
      document.getElementById('cat-label').innerText = '✦ ' + data.category + '  |  SLIDE ' + data.slide_num + ' OF 7';
      document.getElementById('slide-title').innerText = data.title;
      document.getElementById('slide-sub').innerText = data.subtitle;

      // Bullets
      const bulletBox = document.getElementById('bullet-container');
      bulletBox.innerHTML = '';
      data.bullets.forEach(b => {{
        const li = document.createElement('li');
        li.className = 'bullet-item';
        li.innerHTML = '<span class="bullet-dot"></span><span>' + b + '</span>';
        bulletBox.appendChild(li);
      }});

      // Visual items
      document.getElementById('visual-title').innerText = data.visual_label;
      const visBox = document.getElementById('visual-container');
      visBox.innerHTML = '';
      data.visual_items.forEach(([vh, vd]) => {{
        const div = document.createElement('div');
        div.className = 'visual-node';
        div.innerHTML = '<div class="visual-head">' + vh + '</div><div class="visual-desc">' + vd + '</div>';
        visBox.appendChild(div);
      }});

      // Notes
      document.getElementById('notes-content').innerText = data.speaker_note;
      document.getElementById('notes-drawer').style.display = showNotes ? 'block' : 'none';

      // Controls & Progress
      document.getElementById('counter-label').innerText = 'Slide ' + (currentIdx + 1) + ' of ' + deck.length;
      document.getElementById('prog-fill').style.width = Math.round(((currentIdx + 1) / deck.length) * 100) + '%';
      document.getElementById('btn-prev').disabled = (currentIdx === 0);
      document.getElementById('btn-next').disabled = (currentIdx === deck.length - 1);
    }}

    function changeSlide(delta) {{
      const target = currentIdx + delta;
      if (target >= 0 && target < deck.length) {{
        currentIdx = target;
        renderSlide();
      }}
    }}

    function toggleNotes() {{
      showNotes = !showNotes;
      document.getElementById('notes-drawer').style.display = showNotes ? 'block' : 'none';
    }}

    // Keyboard Arrow navigation
    window.addEventListener('keydown', (e) => {{
      if (e.key === 'ArrowRight' || e.key === ' ' || e.key === 'PageDown') {{
        changeSlide(1);
      }} else if (e.key === 'ArrowLeft' || e.key === 'PageUp') {{
        changeSlide(-1);
      }} else if (e.key === 's' || e.key === 'S') {{
        toggleNotes();
      }}
    }});

    renderSlide();
  </script>
</body>
</html>
"""
    with open(HTML_FILE, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"✅ Generated Interactive HTML Presentation: {HTML_FILE}")


if __name__ == "__main__":
    create_pptx_deck()
    create_html_presentation()
